"""Build Dashboard.xlsx: an offline, interactive Excel companion to the web dashboard.

The website is the main product. This workbook exists so that a judge can open a listed dashboard format
(.xlsx) with a double-click and no internet: pick a year and a state, and every figure, the ranking and the
what-if simulator recalculate from live formulas over the same clean data the website uses.

    python -m pipeline.make_excel_dashboard            # -> submission/Dashboard.xlsx
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core import metrics as M

ROOT = Path(__file__).resolve().parents[1]
CLEAN, WEB = ROOT / "data" / "clean", ROOT / "web" / "public" / "data"
OUT = ROOT / "submission"
YEARS = [2023, 2024, 2025]
URL = "https://jomjauh.vercel.app"

INK, MUTED = "1F2937", "6B7280"
RED, BLUE, AMBER, GREEN = "E66767", "3987E5", "E0A030", "199E70"          # problem, opportunity, obstacle, payoff: same as the website
FILL = {k: PatternFill("solid", fgColor=v) for k, v in {"head": "E8EEFB", "input": "FFF7E0", "card": "F5F7FA", "red": "FBE3E3", "blue": "E1EDFC", "amber": "FBF0DC", "green": "DDF3EA"}.items()}
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(size=11, bold=False, color=INK, italic=False):
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)


def put(ws, ref, value, *, size=11, bold=False, color=INK, fill=None, fmt=None, align=None, italic=False, border=False, wrap=False):
    c = ws[ref]
    c.value = value
    c.font = font(size, bold, color, italic)
    if fill:
        c.fill = FILL[fill]
    if fmt:
        c.number_format = fmt
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center" if not wrap else "top", wrap_text=wrap)
    if border:
        c.border = BOX
    return c


# ------------------------------------------------------------------ data
def build_data() -> pd.DataFrame:
    panel = pd.read_parquet(CLEAN / "state_panel.parquet")
    frames = []
    for y in YEARS:
        p = panel[panel.year == y].set_index("code")
        g, b = M.gap_table(p), M.bottlenecks(p)
        d = p.join(g[["potential", "actual", "gap", "gap_rank"]]).join(b[["Access", "Awareness", "Amenities", "bottleneck", "bottleneck_score"]])
        d["bottleneck_shown"] = d.apply(lambda r: r.bottleneck if r.bottleneck_score < 50 else "None", axis=1)
        frames.append(d.reset_index())
    d = pd.concat(frames, ignore_index=True)
    cols = ["label", "year", "code", "state", "region", "visitors_k", "population_k", "visitor_share_pct", "visitors_per_resident", "spend_per_visitor_rm", "receipts_rm_m",
            "avg_length_of_stay", "overnight_share", "paid_accommodation_share", "rooms", "hotels", "occupancy_pct", "experience_score", "potential", "actual", "gap", "gap_rank",
            "Access", "Awareness", "Amenities", "bottleneck_shown", "spend_year"]
    return d[cols]


DESCRIBE = {
    "label": "State (short name)", "year": "Base year", "code": "State code", "state": "State (official name)", "region": "Peninsular / Borneo",
    "visitors_k": "Domestic visitors, thousands (DOSM Domestic Tourism Survey)", "population_k": "Population, thousands (DOSM)", "visitor_share_pct": "Share of all state visits, %",
    "visitors_per_resident": "Visitors per resident", "spend_per_visitor_rm": "Spend per visitor, RM (DOSM, by state, 2023 - carried forward for later years)",
    "receipts_rm_m": "Visitor spending, RM million (estimate after 2023)", "avg_length_of_stay": "Average nights per trip (DOSM 2023)", "overnight_share": "Share of visitors who stay overnight",
    "paid_accommodation_share": "Share of overnight visitors in paid accommodation", "rooms": "Hotel rooms (Tourism Malaysia)", "hotels": "Hotels (Tourism Malaysia)",
    "occupancy_pct": "Average hotel occupancy, % (Tourism Malaysia)", "experience_score": "JomRasa Experience Score 0-100 (public travel text, an indicator not an official statistic)",
    "potential": "What the state can offer, 0-100", "actual": "How visited it already is, 0-100", "gap": "Opportunity (Gap) score = potential - actual", "gap_rank": "Rank by opportunity, 1 = most untapped",
    "Access": "Pillar score, 50 = typical state", "Awareness": "Pillar score, 50 = typical state", "Amenities": "Pillar score, 50 = typical state",
    "bottleneck_shown": "Weakest pillar if below 50, else None", "spend_year": "Year the spending figures come from",
}


def sheet_data(wb: Workbook, d: pd.DataFrame) -> dict[str, str]:
    """Data sheet. Column A is a lookup key (state|year); B is a rank key (rank|year). Returns {column name: letter}."""
    ws = wb.create_sheet("Data")
    names = ["key", "rank_key", *d.columns]
    for j, n in enumerate(names, start=1):
        put(ws, f"{get_column_letter(j)}1", n, bold=True, fill="head", border=True)
        put(ws, f"{get_column_letter(j)}2", DESCRIBE.get(n, "lookup key (do not edit)"), size=9, color=MUTED, italic=True, wrap=True)
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.row_dimensions[2].height = 62
    for i, r in enumerate(d.itertuples(index=False), start=3):
        ws.cell(i, 1, f"{r.label}|{r.year}")
        ws.cell(i, 2, f"{int(r.gap_rank)}|{r.year}")
        for j, v in enumerate(r, start=3):
            ws.cell(i, j, None if pd.isna(v) else (v.item() if hasattr(v, "item") else v))
    ws.freeze_panes = "C3"
    return {n: get_column_letter(j) for j, n in enumerate(names, start=1)}


def look(col: dict[str, str], field: str, key: str) -> str:
    """INDEX/MATCH of a Data column by a key expression."""
    return f"INDEX(Data!${col[field]}$3:${col[field]}$60,MATCH({key},Data!$A$3:$A$60,0))"


# ------------------------------------------------------------------ dashboard sheet
def sheet_dashboard(wb: Workbook, col: dict[str, str], labels: list[str], gini: dict[str, float]) -> None:
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    for c, w in zip("ABCDEFGHIJKLMN", [2, 22, 16, 3, 22, 16, 3, 5, 18, 12, 3, 14, 12, 18]):
        ws.column_dimensions[c].width = w
    put(ws, "B1", "JomJauh - Where should the next visitor go?", size=18, bold=True)
    put(ws, "B2", f"Offline Excel companion of {URL}. Change the two yellow cells; everything else recalculates.", color=MUTED)
    put(ws, "B4", "Year", bold=True); put(ws, "C4", 2025, fill="input", border=True, bold=True, align="center")
    put(ws, "E4", "State", bold=True); put(ws, "F4", "Johor", fill="input", border=True, bold=True, align="center")
    dv = DataValidation(type="list", formula1='"2023,2024,2025"', allow_blank=False); ws.add_data_validation(dv); dv.add("C4")
    dv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(labels) + 1}", allow_blank=False); ws.add_data_validation(dv); dv.add("F4")

    # helper table (columns I:N): the 16 states in opportunity order for the chosen year, with live rank-based Gini
    put(ws, "I6", "States ranked by opportunity", bold=True, size=12)
    for j, h in enumerate(["#", "State", "Opportunity", "", "Visitors (k)", "Rank (asc)"]):
        put(ws, f"{'HIJKLM'[j]}7", h, bold=True, fill="head", border=True, align="center")
    for k in range(1, 17):
        r = 7 + k
        put(ws, f"H{r}", k, align="center", border=True)
        put(ws, f"I{r}", f'=INDEX(Data!${col["label"]}$3:${col["label"]}$60,MATCH({k}&"|"&$C$4,Data!$B$3:$B$60,0))', border=True)
        put(ws, f"J{r}", f'=INDEX(Data!${col["gap"]}$3:${col["gap"]}$60,MATCH({k}&"|"&$C$4,Data!$B$3:$B$60,0))', fmt="+0;-0;0", border=True, align="center")
        put(ws, f"L{r}", f'=INDEX(Data!${col["visitors_k"]}$3:${col["visitors_k"]}$60,MATCH({k}&"|"&$C$4,Data!$B$3:$B$60,0))', fmt="#,##0", border=True)
        put(ws, f"M{r}", f"=RANK(L{r},$L$8:$L$23,1)+COUNTIF($L$8:L{r},L{r})-1", border=True, align="center")
    ws.conditional_formatting.add("H8:J23", FormulaRule(formula=["$I8=$F$4"], fill=FILL["blue"], font=font(bold=True)))

    # the four-step story, national level
    put(ws, "B6", "The story for the chosen year", bold=True, size=12)
    cards = [
        ("B7", "1  THE PROBLEM", "red", "=SUMPRODUCT(($M$8:$M$23>=14)*$L$8:$L$23)/SUM($L$8:$L$23)", "0%", "of all visits go to just 3 of 16 states"),
        ("E7", "2  THE OPPORTUNITY", "blue", "=SUMPRODUCT(($H$8:$H$23<=5)*$N$8:$N$23)/1000", '"+"0.0"M"', "more visitors fit in the 5 most under-visited states"),
        ("B12", "3  THE OBSTACLE", "amber", f'=COUNTIFS(Data!${col["year"]}$3:${col["year"]}$60,$C$4,Data!${col["bottleneck_shown"]}$3:${col["bottleneck_shown"]}$60,"<>None")', '0" of 16 states"', "are held back by one weak point"),
        ("E12", "4  CONCENTRATION", "green", "=SUMPRODUCT((2*$M$8:$M$23-17)*$L$8:$L$23)/(16*SUM($L$8:$L$23))", "0.000", "Gini: 0 = spread evenly, 1 = all in one state"),
    ]
    for ref, title, tone, formula, fmt, line in cards:
        c0, r0 = ref[0], int(ref[1:])
        c1 = chr(ord(c0) + 1)
        for rr in range(r0, r0 + 4):                       # colour first: merged cells cannot be styled afterwards
            for cc in (c0, c1):
                ws[f"{cc}{rr}"].fill = FILL[tone]
        put(ws, f"{c0}{r0}", title, bold=True, size=9, color=MUTED, fill=tone)
        put(ws, f"{c0}{r0 + 1}", formula, size=22, bold=True, fill=tone, fmt=fmt, align="left")
        put(ws, f"{c0}{r0 + 3}", line, size=9, color=MUTED, fill=tone)
        ws.merge_cells(f"{c0}{r0}:{c1}{r0}"); ws.merge_cells(f"{c0}{r0 + 1}:{c1}{r0 + 2}"); ws.merge_cells(f"{c0}{r0 + 3}:{c1}{r0 + 3}")
    # N column: room to grow per ranked state (live, uses the Simulator's assumptions)
    put(ws, "N7", "Room to grow (k)", bold=True, fill="head", border=True, align="center")
    for k in range(1, 17):
        r = 7 + k
        key = f'$I{r}&"|"&$C$4'
        per = f"({look(col, 'overnight_share', key)}*{look(col, 'paid_accommodation_share', key)}*{look(col, 'avg_length_of_stay', key)}/Simulator!$C$9)"
        put(ws, f"N{r}", f"={look(col, 'rooms', key)}*365*MAX(Simulator!$C$8-{look(col, 'occupancy_pct', key)},0)/100/{per}/1000", fmt="#,##0", border=True)

    # the chosen state
    key = '$F$4&"|"&$C$4'
    put(ws, "B17", '=$F$4&" in "&$C$4', bold=True, size=14)
    rows = [
        ("Visitors (thousands)", look(col, "visitors_k", key), "#,##0"), ("Share of all visits", f"{look(col, 'visitor_share_pct', key)}/100", "0.0%"),
        ("Opportunity score", look(col, "gap", key), "+0;-0;0"), ("Rank for opportunity (1 = most untapped)", look(col, "gap_rank", key), '0" of 16"'),
        ("What it can offer (0-100)", look(col, "potential", key), "0"), ("How visited it is (0-100)", look(col, "actual", key), "0"),
        ("Hotels full, on average", f"{look(col, 'occupancy_pct', key)}/100", "0%"), ("Hotel rooms", look(col, "rooms", key), "#,##0"),
        ("Room to grow (extra visitors, thousands)", "INDEX($N$8:$N$23,MATCH($F$4,$I$8:$I$23,0))", "#,##0"),
        ("Spend per visitor, whole trip (RM)", look(col, "spend_per_visitor_rm", key), "#,##0"), ("Travellers rate it (0-100)", look(col, "experience_score", key), "0"),
        ("Access (50 = typical state)", look(col, "Access", key), "0"), ("Awareness (50 = typical state)", look(col, "Awareness", key), "0"), ("Amenities (50 = typical state)", look(col, "Amenities", key), "0"),
        ("What holds it back", look(col, "bottleneck_shown", key), "@"),
    ]
    for i, (label, formula, fmt) in enumerate(rows):
        r = 18 + i
        put(ws, f"B{r}", label, fill="card", border=True)
        ws.merge_cells(f"B{r}:E{r}")
        put(ws, f"F{r}", f"={formula}", fmt=fmt, bold=True, border=True, align="right", fill="amber" if "Access" in label or "Awareness" in label or "Amenities" in label or "holds" in label else None)
    put(ws, "B34", "Spending after 2023 is an estimate: visitors x 2023 spend per visitor (the latest state-level release). Experience scores come from public travel text, not a survey.", size=9, color=MUTED, italic=True)

    chart = BarChart(); chart.type = "bar"; chart.style = 10; chart.title = "Opportunity by state (positive = under-visited)"; chart.legend = None; chart.height, chart.width = 9.5, 15
    chart.add_data(Reference(ws, min_col=10, min_row=7, max_row=23), titles_from_data=True); chart.set_categories(Reference(ws, min_col=9, min_row=8, max_row=23))
    chart.y_axis.scaling.orientation = "minMax"; chart.x_axis.scaling.orientation = "maxMin"; chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].invertIfNegative = False; chart.x_axis.delete = False; chart.y_axis.delete = False; chart.x_axis.tickLblPos = "low"; chart.gapWidth = 60
    ws.add_chart(chart, "H25")

    tr = wb["Trend"]
    line = LineChart(); line.title = "How lopsided tourism has been (Gini)"; line.legend = None; line.height, line.width = 7.5, 15; line.style = 12
    line.add_data(Reference(tr, min_col=2, min_row=1, max_row=len(gini) + 1), titles_from_data=True); line.set_categories(Reference(tr, min_col=1, min_row=2, max_row=len(gini) + 1))
    line.series[0].graphicalProperties.line.solidFill = RED; line.y_axis.number_format = "0.00"; line.x_axis.delete = False; line.y_axis.delete = False; line.series[0].smooth = False
    ws.add_chart(line, "B36")


# ------------------------------------------------------------------ simulator sheet
def sheet_simulator(wb: Workbook, col: dict[str, str], labels: list[str]) -> None:
    ws = wb.create_sheet("Simulator")
    ws.sheet_view.showGridLines = False
    for c, w in zip("ABCDEF", [2, 44, 18, 3, 44, 18]):
        ws.column_dimensions[c].width = w
    put(ws, "B1", "What if some visitors went somewhere quieter?", size=18, bold=True)
    put(ws, "B2", "Change the yellow cells. The year comes from the Dashboard sheet. (The website also handles several destinations at once.)", color=MUTED)
    inputs = [("Take visitors from", "Selangor", "list"), ("Send them to", "Terengganu", "list"), ("Share of its visitors to move", 0.05, "0%"), (None, None, None),
              ("ASSUMPTION: hotels count as full at (occupancy %)", 75, "0"), ("ASSUMPTION: guests per room", 2, "0.0")]
    for i, (label, value, kind) in enumerate(inputs):
        r = 4 + i
        if label is None:
            continue
        put(ws, f"B{r}", label, bold=True)
        put(ws, f"C{r}", value, fill="input", border=True, bold=True, align="center", fmt=None if kind == "list" else kind)
    for ref in ("C4", "C5"):
        dv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(labels) + 1}", allow_blank=False); ws.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="0.3", allow_blank=False); ws.add_data_validation(dv); dv.add("C6")

    o, d = '$C$4&"|"&Dashboard!$C$4', '$C$5&"|"&Dashboard!$C$4'
    per = f"({look(col, 'overnight_share', d)}*{look(col, 'paid_accommodation_share', d)}*{look(col, 'avg_length_of_stay', d)}/$C$9)"
    calc = [
        ("Visitors you asked to move (thousands)", f"={look(col, 'visitors_k', o)}*$C$6", "#,##0"),
        ("Hotel room-nights one visitor needs at the destination", f"={per}", "0.00"),
        ("Spare room-nights at the destination before hotels are full", f"={look(col, 'rooms', d)}*365*MAX($C$8-{look(col, 'occupancy_pct', d)},0)/100", "#,##0"),
        ("Capacity limit: most extra visitors it can take (thousands)", "=C14/C13/1000", "#,##0"),
        ("Visitors actually moved (thousands)", "=IF($C$4=$C$5,0,MIN(C12,C15))", "#,##0"),
        ("Spending gained by the destination (RM million)", f"=C16*{look(col, 'spend_per_visitor_rm', d)}/1000", "#,##0"),
        ("Spending given up by the origin (RM million)", f"=C16*{look(col, 'spend_per_visitor_rm', o)}/1000", "#,##0"),
        ("Net effect for Malaysia (RM million)", "=C17-C18", "+#,##0;-#,##0;0"),
        ("Destination hotel occupancy: before", f"={look(col, 'occupancy_pct', d)}/100", "0.0%"),
        ("Destination hotel occupancy: after", f"=C20+C16*1000*C13/({look(col, 'rooms', d)}*365)", "0.0%"),
    ]
    put(ws, "B11", "Result", bold=True, size=12)
    for i, (label, formula, fmt) in enumerate(calc):
        r = 12 + i
        put(ws, f"B{r}", label, fill="card", border=True)
        put(ws, f"C{r}", formula, fmt=fmt, bold=True, border=True, align="right", fill="green" if i in (4, 7) else None)
    put(ws, "B23", '=IF($C$4=$C$5,"Pick two different states.",IF(C12>C15,"Capacity limit reached: only "&TEXT(C16,"#,##0")&"k of "&TEXT(C12,"#,##0")&"k fit. Here the fix is more rooms, not more marketing.","The destination can absorb all of them."))', bold=True, color=GREEN)
    put(ws, "B24", "This is rebalancing, not new money: the same visitors spend in a different place, so the national total barely moves.", size=9, color=MUTED, italic=True)

    # concentration before / after, by ranking the 16 states again after the move
    put(ws, "E11", "Is tourism more evenly spread afterwards?", bold=True, size=12)
    for j, h in enumerate(["State", "Visitors after (k)"]):
        put(ws, f"{'EF'[j]}12", h, bold=True, fill="head", border=True)
    for k, name in enumerate(labels):
        r = 13 + k
        put(ws, f"E{r}", name, border=True)
        put(ws, f"F{r}", f'={look(col, "visitors_k", f"$E{r}&" + chr(34) + "|" + chr(34) + "&Dashboard!$C$4")}-IF($E{r}=$C$4,$C$16,0)+IF($E{r}=$C$5,$C$16,0)', fmt="#,##0", border=True)
        ws[f"G{r}"] = f"=RANK(F{r},$F$13:$F$28,1)+COUNTIF($F$13:F{r},F{r})-1"
        ws[f"G{r}"].font = font(9, color="BBBBBB")
    put(ws, "E30", "Concentration (Gini) before", fill="card", border=True); put(ws, "F30", "=Dashboard!E13", fmt="0.0000", bold=True, border=True)
    put(ws, "E31", "Concentration (Gini) after", fill="card", border=True); put(ws, "F31", "=SUMPRODUCT((2*$G$13:$G$28-17)*$F$13:$F$28)/(16*SUM($F$13:$F$28))", fmt="0.0000", bold=True, border=True, fill="green")
    put(ws, "E32", "Lower is more even.", size=9, color=MUTED, italic=True)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    d = build_data()
    labels = sorted(d.label.unique())
    trend = json.loads((WEB / "trend.json").read_text(encoding="utf8"))
    jr = json.loads((WEB / "jomrasa.json").read_text(encoding="utf8"))

    wb = Workbook()
    start = wb.active; start.title = "Start here"; start.sheet_view.showGridLines = False
    start.column_dimensions["B"].width = 110
    lines = [
        ("JomJauh + JomRasa - Team ANAK SUNWAY - DOSM Datathon 2026", 18, True, INK),
        (f"The full interactive dashboard is a website: {URL}  (no login needed). This workbook is its offline companion.", 11, False, INK), ("", 11, False, INK),
        ("How to use this workbook", 13, True, INK),
        ("1. Open the Dashboard sheet. Pick a YEAR and a STATE in the two yellow cells. Every number, the ranking and the chart recalculate.", 11, False, INK),
        ("2. Open the Simulator sheet. Choose where visitors come from, where they go and how many. It shows the money moved, whether the hotels can cope, and whether tourism ends up more evenly spread.", 11, False, INK),
        ("3. The Data, Trend, Travellers and Sources sheets hold the clean data behind everything, with a description under every column name.", 11, False, INK), ("", 11, False, INK),
        ("The four colours mean the same thing as on the website: red = the problem (crowding), blue = the opportunity, amber = the obstacle, green = the payoff.", 11, False, MUTED),
        ("Nothing here is typed in by hand: figures come from DOSM, Tourism Malaysia, data.gov.my and OpenStreetMap through the project's pipeline. No macros - it is safe to open.", 11, False, MUTED),
        ("Not in this workbook (website only): the interactive map, adjustable index weights and sensitivity test, several destinations in the simulator, the hover guide, and the JomRasa trip planner.", 11, False, MUTED),
    ]
    for i, (text, size, bold, color) in enumerate(lines, start=2):
        put(start, f"B{i}", text, size=size, bold=bold, color=color, wrap=True)
        start.row_dimensions[i].height = 34 if len(text) > 110 else 22

    lists = wb.create_sheet("Lists"); put(lists, "A1", "States", bold=True)
    for i, n in enumerate(labels, start=2):
        lists[f"A{i}"] = n

    tr = wb.create_sheet("Trend")
    put(tr, "A1", "Year", bold=True, fill="head"); put(tr, "B1", "Gini of visitors across states", bold=True, fill="head")
    for i, (y, g) in enumerate(sorted(trend["gini_by_year"].items()), start=2):
        tr[f"A{i}"], tr[f"B{i}"] = int(y), g
    years = sorted({int(y) for v in trend["visitors"].values() for y in v})
    put(tr, "D1", "Visitors by state and year (thousands)", bold=True, fill="head")
    for j, y in enumerate(years, start=5):
        put(tr, f"{get_column_letter(j)}1", y, bold=True, fill="head")
    names = d.drop_duplicates("code").set_index("code").label.to_dict()
    for i, (code, series) in enumerate(sorted(trend["visitors"].items()), start=2):
        tr[f"D{i}"] = names.get(code, code)
        for j, y in enumerate(years, start=5):
            tr.cell(i, j, series.get(str(y)))
    tr.column_dimensions["B"].width = 30; tr.column_dimensions["D"].width = 22

    col = sheet_data(wb, d)
    sheet_dashboard(wb, col, labels, trend["gini_by_year"])
    sheet_simulator(wb, col, labels)

    tv = wb.create_sheet("Travellers")
    put(tv, "A1", "What travellers say, by state and topic (JomRasa). Sentiment 0-100, 50 = neutral; n = number of mentions. An indicator from public travel text, not an official statistic.", bold=True)
    topics = pd.DataFrame(jr["topics"])
    pivot = topics.pivot_table(index="code", columns="topic", values="sentiment").round(0)
    put(tv, "A3", "State", bold=True, fill="head")
    for j, t in enumerate(pivot.columns, start=2):
        put(tv, f"{get_column_letter(j)}3", t.replace("_", " "), bold=True, fill="head"); tv.column_dimensions[get_column_letter(j)].width = 17
    for i, (code, r) in enumerate(pivot.iterrows(), start=4):
        tv[f"A{i}"] = names.get(code, code)
        for j, v in enumerate(r, start=2):
            tv.cell(i, j, None if pd.isna(v) else float(v))
    tv.column_dimensions["A"].width = 20

    src = wb.create_sheet("Sources"); src.column_dimensions["A"].width = 46; src.column_dimensions["B"].width = 100
    put(src, "A1", "Source", bold=True, fill="head"); put(src, "B1", "What it provides", bold=True, fill="head")
    for i, (a, b) in enumerate([
        ("DOSM Domestic Tourism Survey 2023-2025 (OpenDOSM)", "Visitors by state, tourists / day-trippers, origin-destination flows; spending and length of stay by state (2023)."),
        ("Tourism Malaysia Paid Accommodation Survey", "Hotels, rooms, occupancy and hotel guests by state, year and quarter."),
        ("data.gov.my", "Population, household income, poverty, basic amenities and GDP by state."),
        ("OpenStreetMap (ODbL)", "Attractions, nature and heritage sites, airports, rail stations, hotels and eateries."),
        ("geoBoundaries", "State boundaries for the map."),
        ("Public travel text (Exa search, YouTube Data API)", "De-identified travel posts and comments in Malay, English and Mandarin, tagged by a language model (JomRasa). No usernames are stored."),
        ("Wikivoyage (CC BY-SA 4.0), Open-Meteo / ERA5 (CC BY 4.0)", "Traveller guide: places, eateries, places to sleep; ten years of rainfall for the best months to visit."),
        ("Full method, data dictionary and source code", f"{URL}/methodology  and  https://github.com/ZenBen5173/jomjauh-jomrasa"),
    ], start=2):
        put(src, f"A{i}", a, wrap=True); put(src, f"B{i}", b, wrap=True); src.row_dimensions[i].height = 32

    wb._sheets = [wb[n] for n in ["Start here", "Dashboard", "Simulator", "Data", "Trend", "Travellers", "Sources", "Lists"]]
    wb["Lists"].sheet_state = "hidden"
    wb.active = 1
    path = OUT / "Dashboard.xlsx"
    wb.save(path)
    print(f"wrote {path} ({path.stat().st_size // 1024} KB), {len(d)} state-year rows")


if __name__ == "__main__":
    main()
